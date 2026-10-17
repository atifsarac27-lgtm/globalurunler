"""birlestir.py testleri. Ağ kullanmaz."""
import os
import tempfile
import unittest

from openpyxl import load_workbook

import birlestir as b


VERI = {"tum": [
    {"sinif": "ulke", "marka": "KFC", "ulke": "Japonya", "ad": "Yeni Go Bucket",
     "kaynaklar": [{"adres": "https://kfc.jp/x"}], "tespit_tarihi": "2026-10-13"},
    {"sinif": "dunya", "marka": "KFC", "ulke": "Dünya", "ad": "Global menü", "kaynaklar": []},
    {"sinif": "ulke", "marka": "Popeyes", "ulke": "Türkiye", "ad": "Honey Biscuit",
     "site_adresi": "https://popeyes.tr/", "tespit_tarihi": ""},
]}
HABER = {"yeni_urun": [
    {"metin": "Popeyes launches new biscuit", "url": "https://n/1", "marka": "Popeyes"},
]}


class BirlestirTesti(unittest.TestCase):
    def test_resmi_yalnizca_ulke_sinifi(self):
        satirlar = b.resmi_satirlari(VERI)
        self.assertEqual(len(satirlar), 2)
        self.assertTrue(all(s[3] == "Resmi site" for s in satirlar))

    def test_resmi_link_yoksa_site_adresi_kullanilir(self):
        satirlar = b.resmi_satirlari(VERI)
        popeyes = [s for s in satirlar if s[0] == "Popeyes"][0]
        self.assertEqual(popeyes[4], "https://popeyes.tr/")

    def test_haber_satirlari_yeni_urun_alanini_okur(self):
        satirlar = b.haber_satirlari(HABER)
        self.assertEqual(len(satirlar), 1)
        self.assertEqual(satirlar[0][3], "Haber")

    def test_excel_yazilir_ve_ozet_sayfasi_var(self):
        with tempfile.TemporaryDirectory() as d:
            yol = os.path.join(d, "birlesik.xlsx")
            b.yaz(yol, b.resmi_satirlari(VERI) + b.haber_satirlari(HABER))
            wb = load_workbook(yol)
            self.assertEqual(wb.sheetnames, ["Yeni Ürünler", "Özet"])
            self.assertEqual(wb["Yeni Ürünler"].max_row, 1 + 3)


if __name__ == "__main__":
    unittest.main()
