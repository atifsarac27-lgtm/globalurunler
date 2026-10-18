"""Resmi site sonuçlarını (radar veri.json) haber sonuçlarıyla tek Excel'de birleştirir.

Kullanım:
    python birlestir.py --veri rapor/veri.json --haber rapor/haber_urunleri.json --cikti rapor/birlesik.xlsx

Kurallar:
- Resmi site: radar raporuyla aynı kural (sinif == "ulke" ürünler).
- Haber: yeni_urun_siniflandir.py çıktısında yeni_urun = true olanlar.
- Aynı ürün iki kaynakta da geçebilir; bu dosya otomatik eşleştirme yapmaz, "Kaynak türü" sütunu ile ayrılır.
- Haber dosyası yoksa (haber adımı başarısız olduysa) yalnızca resmi siteler yazılır.
"""
import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASLIKLAR = ["Marka", "Ülke", "Ürün / Başlık", "Kaynak türü", "Kaynak link", "Tespit tarihi"]
FONT = "Arial"


def resmi_satirlari(veri):
    satirlar = []
    for a in veri.get("tum", []):
        if a.get("sinif") != "ulke":
            continue
        kaynak = (a.get("kaynaklar") or [{}])[0].get("adres") or a.get("site_adresi") or ""
        satirlar.append([a.get("marka", ""), a.get("ulke", ""), a.get("ad", ""),
                         "Resmi site", kaynak, a.get("tespit_tarihi") or ""])
    return satirlar


def haber_satirlari(haber):
    return [[k.get("marka_site") or k.get("marka", ""), k.get("ulke") or "-", k.get("metin", ""),
             "Haber", k.get("url", ""), k.get("tarih", "")]
            for k in haber.get("yeni_urun", [])]


def yaz(cikti, satirlar):
    wb = Workbook()
    ws = wb.active
    ws.title = "Yeni Ürünler"
    ws.append(BASLIKLAR)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="1F3864")
    for s in sorted(satirlar, key=lambda s: (s[0], s[1], s[2])):
        ws.append(s)
        r = ws.max_row
        for col in range(1, 7):
            ws.cell(r, col).font = Font(name=FONT, size=10)
            ws.cell(r, col).alignment = Alignment(vertical="top", wrap_text=(col == 3))
    for col, w in zip("ABCDEF", [16, 14, 90, 14, 60, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"

    ozet = wb.create_sheet("Özet")
    ozet.append(["Marka", "Resmi site", "Haber", "Toplam"])
    for c in ozet[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="1F3864")
    markalar = sorted({s[0] for s in satirlar})
    for m in markalar:
        r = ozet.max_row + 1
        ozet.append([
            m,
            f"=COUNTIFS('Yeni Ürünler'!$A:$A,A{r},'Yeni Ürünler'!$D:$D,\"Resmi site\")",
            f"=COUNTIFS('Yeni Ürünler'!$A:$A,A{r},'Yeni Ürünler'!$D:$D,\"Haber\")",
            f"=B{r}+C{r}",
        ])
    if markalar:
        son = ozet.max_row
        ozet.append(["Toplam", f"=SUM(B2:B{son})", f"=SUM(C2:C{son})", f"=SUM(D2:D{son})"])
    for row in ozet.iter_rows(min_row=2, max_row=ozet.max_row):
        for c in row:
            c.font = Font(name=FONT, size=10, bold=(c.row == ozet.max_row))
    ozet.column_dimensions["A"].width = 18
    for col in "BCD":
        ozet.column_dimensions[col].width = 14

    wb.save(cikti)


def main():
    p = argparse.ArgumentParser(description="Resmi site ve haber sonuçlarını birleştir")
    p.add_argument("--veri", required=True, help="radar veri.json yolu")
    p.add_argument("--haber", required=True, help="yeni_urun_siniflandir.py çıktısı")
    p.add_argument("--cikti", required=True, help="birleşik Excel yolu")
    a = p.parse_args()

    with open(a.veri, encoding="utf-8") as f:
        veri = json.load(f)
    if os.path.exists(a.haber):
        with open(a.haber, encoding="utf-8") as f:
            haber = json.load(f)
    else:
        print(f"UYARI: {a.haber} yok, yalnızca resmi siteler yazılacak")
        haber = {"yeni_urun": []}

    resmi = resmi_satirlari(veri)
    haberler = haber_satirlari(haber)
    yaz(a.cikti, resmi + haberler)
    print(f"Resmi site: {len(resmi)}, haber: {len(haberler)}, toplam: {len(resmi) + len(haberler)}")
    print(f"Çıktı: {a.cikti}")


if __name__ == "__main__":
    main()
